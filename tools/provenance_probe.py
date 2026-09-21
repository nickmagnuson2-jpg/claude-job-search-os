#!/usr/bin/env python3
"""THE PROVENANCE PROBE: which printed numbers have no formula behind them?

THE RULE BEING AUDITED. Every derived number a reader might act on must be traceable to a
formula over the client's own rows. A number that exists only in a Python dict, or is pasted
as a typed value into a sheet, fails that. The question has one right answer for every
engagement, which is why the instrument lives here and the sheet names, exemptions and file
paths stay with the caller.

HOW IT WORKS. Two passes over a RECALCULATED copy of the workbook, so values are real:
  1. Index every cell holding a FORMULA, paired with its computed value.
  2. Walk every numeric leaf in the analysis JSON and look for a formula cell whose value
     matches. Report the ones with no match.

WHAT A MISS MEANS. Either the number is genuinely unbacked, or it is backed on a sheet this
audit cannot see. Both need a human decision, which is why the CLI PRINTS rather than asserts
unless `--fail-on-unbacked` is passed. A typed value that happens to equal a formula's result
elsewhere will match, so this is a LOWER BOUND on the problem, never an all-clear.

WHICH SHEETS YOU EXCLUDE IS THE WHOLE CORRECTNESS OF THE AUDIT, and it is the caller's call.
In the engagement this was extracted from, two per-row classification sheets carried ~242,000
seconds-and-hours cells. Indexing them meant roughly ONE INTEGER IN FIVE under 7,000 matched
some cell by coincidence, so numbers were reported "backed" by accident. The first run said 23
unbacked; that was an instrument artifact reported as a measurement. The tell was visible and
ignored: the index jumped from 1,347 cells to 242,193 when those sheets were added, and nobody
asked why. A number is BACKED only if a SUMMARY cell computes it; a per-row intermediate is a
coincidence with the same digits.

So `coincidence_rate()` is part of the instrument, not a nicety. It samples integers the
analysis never printed and reports how often the index matches one anyway. That rate is the
probe's own false-positive floor. Run it, print it beside the verdict, and an index that has
quietly swallowed a per-row sheet announces itself as a number instead of as a comment.

EMPTY INPUT IS AN ERROR, NOT A NEGATIVE RESULT. An index with no formula cells raises, and so
does an analysis with no numeric leaves. A probe that scans nothing and reports "everything is
backed" is the defect this family is named for.

USAGE as a library (the policy -- paths, sheet names, exemptions -- stays with the caller):

    from provenance_probe import load_workbooks, index_formula_cells, trace_numbers

    formulas, values = load_workbooks(WB, RECALC)
    backed = index_formula_cells(formulas, values, exclude_sheets=RAW | PER_ROW)
    report = trace_numbers(json.loads(JSON_IN.read_text()), backed, exempt=EXEMPT)
    print(report.render())

USAGE as a CLI:

    PYTHONIOENCODING=utf-8 python3 tools/provenance_probe.py \\
        --workbook book.xlsx --values recalculated.xlsx --json analysis.json \\
        --exclude-sheet "12 Per-row detail" --exempt "4200=given by the client, not derived"

Promoted 2026-09-21 from an engagement's `scripts/audit_provenance.py`.
"""

from __future__ import annotations

import argparse
import json
import random
import sys
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable, Iterator, Mapping, Sequence

from json_leaves import leaves, numeric

# Match with TOLERANCE, not exact equality. A JSON rounds for readability (451.2) while the
# workbook carries full precision (451.23456789012), and rates are often percentages in the
# JSON against fractions in the workbook. An exact-key matcher reported three of a
# deliverable's headline figures as unbacked when all three were fine, which sends the reader
# to fix the wrong thing. RELATIVE_TOLERANCE handles precision; ABSOLUTE_FLOOR handles small
# integers, where a relative tolerance collapses to nothing.
RELATIVE_TOLERANCE = 1e-4
ABSOLUTE_FLOOR = 0.051


@dataclass(frozen=True)
class BackedCell:
    """One formula cell and the value it computed to."""

    value: float
    ref: str


@dataclass
class ProvenanceReport:
    indexed: int
    traced: int
    unbacked: list[tuple[str, float]]

    @property
    def clean(self) -> bool:
        """True only when every numeric leaf was traced or explicitly exempted."""
        return not self.unbacked

    def render(self, max_per_block: int = 6) -> str:
        lines = [
            f"formula-backed cells indexed: {self.indexed}",
            f"analysis numbers traced:      {self.traced}",
            f"analysis numbers UNBACKED:    {len(self.unbacked)}",
            "",
        ]
        by_block: dict[str, list[tuple[str, float]]] = {}
        for path, val in self.unbacked:
            by_block.setdefault(path.split(".")[0].split("[")[0], []).append((path, val))
        for block, items in sorted(by_block.items(), key=lambda kv: -len(kv[1])):
            lines.append(f"  {block}  ({len(items)} unbacked)")
            for path, val in items[:max_per_block]:
                lines.append(f"      {path:58s} {val:,.2f}")
            if len(items) > max_per_block:
                lines.append(f"      ... and {len(items) - max_per_block} more")
        return "\n".join(lines)


def numeric_leaves(node, path: str = "") -> Iterator[tuple[str, float]]:
    """Every number in a nested JSON structure, with its dotted path.

    The walk and the path convention live in tools/json_leaves.py, shared with the decision
    extractor, which wants STRING leaves out of the same shape. Two copies of a path
    convention drift on the first edge case and then two reports name the same leaf
    differently. Booleans are excluded by `numeric`; see its docstring for why that matters.
    """
    for p, v in leaves(node, path, keep=numeric):
        yield p, float(v)


def load_workbooks(formula_path, values_path=None):
    """Open the workbook twice: once for formula text, once for computed values.

    `values_path` defaults to the same file, which is correct only when that file was
    RECALCULATED by something that evaluates formulas. openpyxl never evaluates, so a workbook
    openpyxl wrote carries no cached values and every lookup returns None -- an index of zero
    backed cells, which `index_formula_cells` raises on rather than reporting as a clean run.
    """
    import openpyxl  # imported here so the pure functions are usable without it

    return (
        openpyxl.load_workbook(formula_path),
        openpyxl.load_workbook(values_path or formula_path, data_only=True),
    )


def index_formula_cells(
    formulas_wb, values_wb, exclude_sheets: Iterable[str] = ()
) -> list[BackedCell]:
    """Index every formula cell that computed to a number, skipping excluded sheets.

    Raises ValueError on an empty index: that is an unrecalculated workbook or an exclusion
    list that swallowed everything, and either way the run measured nothing.
    """
    skip = set(exclude_sheets)
    unknown = skip - set(formulas_wb.sheetnames)
    if unknown:
        raise ValueError(
            "exclude_sheets names sheets that are not in the workbook: "
            + ", ".join(sorted(unknown))
            + ". A misspelled exclusion silently indexes the sheet it meant to skip."
        )
    backed: list[BackedCell] = []
    for name in formulas_wb.sheetnames:
        if name in skip:
            continue
        fs, vs = formulas_wb[name], values_wb[name]
        for row in fs.iter_rows():
            for cell in row:
                v = cell.value
                if not (isinstance(v, str) and v.startswith("=")):
                    continue
                got = vs[cell.coordinate].value
                if isinstance(got, (int, float)) and not isinstance(got, bool):
                    backed.append(BackedCell(float(got), f"{name}!{cell.coordinate}"))
    if not backed:
        raise ValueError(
            "no formula cell computed to a number. The workbook is not recalculated, or "
            "exclude_sheets removed every sheet that has formulas."
        )
    return backed


def find_backing(
    value: float,
    backed: Sequence[BackedCell],
    *,
    also_try_fraction: bool = True,
) -> str | None:
    """The cell reference backing `value`, or None.

    `also_try_fraction` tries value/100 as well, because a rate printed as 12.5 in the JSON is
    0.125 in the workbook. Turn it off when the analysis holds no percentages -- it widens the
    match surface and every widening raises the coincidence rate.

    THE FRACTION PATH IS GATED TO PLAUSIBLE PERCENTAGES (0 < value <= 100), and that gate was
    measured, not assumed. Ungated, 999 traces to a cell holding 10 (999/100 = 9.99, inside the
    absolute floor), and on the engagement this was promoted from the ungated path carried a
    COINCIDENCE FLOOR OF 11.2% against 1.5% gated -- one traced number in nine was a digit
    accident. Gating kept 54 of the 58 numbers that only the fraction path could trace and
    dropped the four that were the 999-shaped kind. A value above 100 is not a percentage; if
    an analysis really prints rates above 100%, match them on the primary path instead.
    """
    targets = [value]
    if also_try_fraction and 0 < value <= 100:
        targets.append(value / 100.0)
    for target in targets:
        tol = max(abs(target) * RELATIVE_TOLERANCE, ABSOLUTE_FLOOR)
        for cell in backed:
            if abs(cell.value - target) <= tol:
                return cell.ref
    return None


def trace_numbers(
    analysis,
    backed: Sequence[BackedCell],
    exempt: Mapping[int, str] | None = None,
    *,
    also_try_fraction: bool = True,
) -> ProvenanceReport:
    """Trace every numeric leaf of `analysis` to a formula cell.

    `exempt` maps an integer to the reason it legitimately has no formula (a constant the
    client stated, a row count). The reason is required by the mapping's shape so that an
    exemption cannot be added without one.

    Raises ValueError when `analysis` holds no numbers: a traced-0-of-0 run is vacuous and
    renders as a clean report.
    """
    exempt = exempt or {}
    traced, unbacked = 0, []
    seen = False
    for path, val in numeric_leaves(analysis):
        seen = True
        if find_backing(val, backed, also_try_fraction=also_try_fraction):
            traced += 1
        elif val == int(val) and int(val) in exempt:
            traced += 1
        else:
            unbacked.append((path, val))
    if not seen:
        raise ValueError("the analysis holds no numbers; there is nothing to trace")
    return ProvenanceReport(indexed=len(backed), traced=traced, unbacked=unbacked)


def coincidence_rate(
    backed: Sequence[BackedCell],
    *,
    samples: int = 400,
    low: int = 1,
    high: int = 7000,
    exclude: Iterable[float] = (),
    seed: int | None = 0,
    also_try_fraction: bool = True,
) -> float:
    """The share of integers the analysis never printed that the index matches anyway.

    This is the probe's own false-positive floor, and it is the measurement that would have
    caught the per-row-sheet artifact at the time instead of a week later. A healthy index sits
    near zero; a rate of 0.2 means one in five "backed" verdicts is a digit coincidence.

    `exclude` is the analysis's own numbers, so a real backing is not counted as a coincidence.
    """
    if not backed:
        raise ValueError("cannot measure the coincidence rate of an empty index")
    if samples <= 0:
        raise ValueError("samples must be positive; a zero-sample rate measures nothing")
    if low > high:
        raise ValueError(f"empty sampling range: low={low} exceeds high={high}")
    rng = random.Random(seed)
    skip = {float(x) for x in exclude}
    hits = drawn = 0
    while drawn < samples:
        candidate = float(rng.randint(low, high))
        if candidate in skip:
            continue
        drawn += 1
        if find_backing(candidate, backed, also_try_fraction=also_try_fraction):
            hits += 1
    return hits / drawn


def _parse_exempt(raw: Sequence[str]) -> dict[int, str]:
    out: dict[int, str] = {}
    for item in raw:
        number, _, reason = item.partition("=")
        if not reason.strip():
            raise SystemExit(
                f"--exempt {item!r} has no reason. An exemption without a written reason is "
                "how an unbacked number becomes permanent."
            )
        out[int(number.strip())] = reason.strip()
    return out


def main(argv: Sequence[str] | None = None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    ap.add_argument("--workbook", required=True, type=Path, help="xlsx holding formula text")
    ap.add_argument(
        "--values",
        type=Path,
        help="RECALCULATED xlsx holding computed values (default: --workbook)",
    )
    ap.add_argument("--json", required=True, type=Path, help="analysis JSON to trace")
    ap.add_argument(
        "--exclude-sheet",
        action="append",
        default=[],
        metavar="NAME",
        help="sheet to keep OUT of the backing index; repeatable. Per-row sheets belong here",
    )
    ap.add_argument(
        "--exempt",
        action="append",
        default=[],
        metavar="N=REASON",
        help="integer that legitimately has no formula, with the reason; repeatable",
    )
    ap.add_argument(
        "--no-fraction-match",
        action="store_true",
        help="do not also try value/100 (use when the analysis holds no percentages)",
    )
    ap.add_argument(
        "--coincidence",
        type=int,
        default=0,
        metavar="N",
        help="also measure the index's false-match rate over N sampled integers",
    )
    ap.add_argument(
        "--fail-on-unbacked",
        action="store_true",
        help="exit 2 when anything is unbacked (default: print and exit 0, because a miss "
        "needs a human decision)",
    )
    a = ap.parse_args(argv)

    formulas, values = load_workbooks(a.workbook, a.values)
    backed = index_formula_cells(formulas, values, exclude_sheets=a.exclude_sheet)
    analysis = json.loads(a.json.read_text())
    fraction = not a.no_fraction_match
    report = trace_numbers(
        analysis,
        backed,
        exempt=_parse_exempt(a.exempt),
        also_try_fraction=fraction,
    )
    print(report.render())
    if a.coincidence:
        rate = coincidence_rate(
            backed,
            samples=a.coincidence,
            exclude=[v for _, v in numeric_leaves(analysis)],
            also_try_fraction=fraction,
        )
        print(f"\nindex coincidence rate: {rate:.1%} over {a.coincidence} sampled integers")
        print("  (the probe's own false-positive floor; a high rate means the index is "
              "matching by digits, not by provenance)")
    return 2 if (a.fail_on_unbacked and not report.clean) else 0


if __name__ == "__main__":
    sys.exit(main())
